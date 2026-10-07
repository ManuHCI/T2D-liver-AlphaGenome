"""
Step 15 - Assemble Supplementary Tables S1-S9 into one Excel workbook.

Reads results/, results_hybrid/, results_prox1/, data/ ; writes supplementary/Supplementary_Tables_S1-S9.xlsx
Run:  python 15_supplementary_tables.py
"""
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT = "supplementary/Supplementary_Tables_S1-S9.xlsx"


def lit_table():
    path = "results/T2D_tree_literature_check.docx"
    if not os.path.exists(path):
        return pd.DataFrame()
    import docx
    t = docx.Document(path).tables[0]
    rows = [[c.text for c in r.cells] for r in t.rows]
    return pd.DataFrame(rows[1:], columns=rows[0])


def main():
    os.makedirs("supplementary", exist_ok=True)
    var = pd.read_csv("data/t2d_variants_for_alphagenome.csv")
    var = var[["rsid", "variant_id", "chrom", "pos", "ref", "alt", "risk_allele", "beta_alt", "pip", "consequence",
               "study_id", "n_studies_supporting", "studies_supporting"]]
    bench = pd.read_csv("results/benchmark_per_variant_scores.csv")
    b_dir = pd.read_csv("results/benchmark_direction.csv")
    b_gene = pd.read_csv("results/benchmark_gene_choice.csv")
    b_rule = pd.read_csv("results/benchmark_rules.csv")
    links = pd.read_csv("results_hybrid/level1_all_nearest_genes.csv")
    genes = pd.read_csv("results_hybrid/level0_genes.csv")
    prio = pd.read_csv("results_hybrid/level1_prioritisation.csv").drop(columns=["Unnamed: 0"], errors="ignore")
    enh = pd.read_csv("results/level0_enhancers.csv")
    enh = enh[enh["enhancer_level"]].sort_values("n_TF_tracks_changed", ascending=False)
    tfb = pd.read_csv("results/level0_tf_binding.csv")
    nodes = pd.read_csv("results_hybrid/tree_nodes.csv")
    nodes = nodes[nodes["level"] >= 2].sort_values("p_push").head(300)
    paths = pd.read_csv("results_hybrid/tree_paths.csv")
    cc = pd.read_csv("results_hybrid/control_comparison.csv")
    gd = pd.read_csv("results_hybrid/genetic_direction_test.csv")
    p = "results_prox1/"
    sheets = [
        ("README", None, None),
        ("S1_T2D_variants", var, "Table S1. 923 fine-mapped non-coding T2D variants (SuSiE PIP >= 0.2) scored with "
                                  "AlphaGenome. risk_allele = ALT if beta_alt > 0."),
        ("S2_liver_eQTL_benchmark", bench, "Table S2. 328 causal GTEx v8 liver eQTLs (SuSiE PIP >= 0.5, eQTL "
                                           "Catalogue QTD000266): eGene, nearest gene, AlphaGenome top genes and "
                                           "AlphaGenome scores (HepG2/liver/hepatocyte RNA-seq)."),
        ("S3_benchmark_summary", None, "Table S3. Benchmark summaries: direction, gene choice, assignment rules."),
        ("S4_hybrid_level1_links", links, "Table S4. Nearest protein-coding gene for every scored T2D variant with "
                                          "AlphaGenome hepatic score; significant = |quantile| >= 0.99 "
                                          "(108 variants kept)."),
        ("S5_level1_genes_priority", None, "Table S5. 91 Level-1 genes and the 7-criterion prioritisation of the "
                                           "108 variant-gene pairs."),
        ("S6_level0_mechanisms", None, "Table S6. Enhancer-level variants (>= 100 HepG2 ChIP-TF tracks change "
                                       "together) and TF-specific binding changes."),
        ("S7_network_tree", None, "Table S7. Network propagation (Levels 2+): top 300 nodes by p, exploratory "
                                  "paths, and validation against three yardsticks."),
        ("S8_PROX1_locus", None, "Table S8. PROX1 locus: LD/allele frequency, tissue specificity, GTEx eQTL, "
                                 "in-silico mutagenesis, motif changes, ENCODE ChIP-seq, colocalisation."),
        ("S9_literature_check", lit_table(), "Table S9. Literature check of 11 tree claims (PubMed)."),
    ]
    blocks = {
        "S3_benchmark_summary": [("Direction of effect", b_dir), ("Target-gene choice", b_gene),
                                 ("Assignment rules (gene AND direction correct)", b_rule)],
        "S5_level1_genes_priority": [("91 Level-1 genes (PIP-weighted score)", genes),
                                     ("Prioritisation of 108 variant-gene pairs", prio)],
        "S6_level0_mechanisms": [("Enhancer-level variants", enh), ("TF-specific binding changes", tfb)],
        "S7_network_tree": [("Validation (GEO, Open Targets)", cc), ("Validation (FinnGen genetic direction)", gd),
                            ("Top 300 network nodes (Levels 2+)", nodes), ("Exploratory paths", paths)],
        "S8_PROX1_locus": [("LD with rs17712208 (1000G EUR, Ensembl)", pd.read_csv(p + "test1_ld.csv")),
                           ("Allele frequencies", pd.read_csv(p + "test1_allele_frequencies.csv")),
                           ("AlphaGenome tissue specificity", pd.read_csv(p + "test2_tissue_specificity.csv")),
                           ("rs17712208 DNase change by biosample",
                            pd.read_csv(p + "test2_rs17712208_dnase_by_biosample.csv")),
                           ("GTEx v8 eQTL (PROX1)", pd.read_csv(p + "test3_gtex_eqtl.csv")),
                           ("In-silico mutagenesis (mean of 3 alternatives per position)",
                            pd.read_csv(p + "test4_ism_by_position.csv")),
                           ("JASPAR 2024 motif changes (ref score >= 0.80)", pd.read_csv(p + "test4_motif_changes.csv")),
                           ("AlphaGenome ChIP-TF changes (HepG2)", pd.read_csv(p + "test4_alphagenome_chip_tf.csv")),
                           ("ENCODE ChIP-seq peaks (selected TFs)",
                            pd.read_csv(p + "encode_selected_tf_peaks.csv").query("bound == True")),
                           ("ENCODE TF clusters (UCSC)", pd.read_csv(p + "encode_tf_clusters.csv")),
                           ("Colocalisation: own CLPP", pd.read_csv(p + "coloc_own_clpp.csv")),
                           ("Colocalisation: Open Targets (traits of interest)",
                            pd.read_csv(p + "coloc_open_targets.csv").dropna(subset=["group"])),
                           ("PheWAS (Open Targets credible sets + GWAS Catalog)", pd.read_csv(p + "test5_phewas.csv"))],
    }
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        readme = pd.DataFrame({"Sheet": [s for s, _, _ in sheets[1:]], "Content": [c for _, _, c in sheets[1:]]})
        readme.to_excel(xw, sheet_name="README", index=False, startrow=3)
        for name, df, cap in sheets[1:]:
            if name in blocks:
                row = 2
                for title, d in blocks[name]:
                    pd.DataFrame([[title]]).to_excel(xw, sheet_name=name, index=False, header=False, startrow=row)
                    d.to_excel(xw, sheet_name=name, index=False, startrow=row + 1)
                    row += len(d) + 4
            else:
                df.to_excel(xw, sheet_name=name, index=False, startrow=2)
    wb = load_workbook(OUT)
    caps = {s: c for s, _, c in sheets}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=9, bold=c.font.bold)
        if ws.title == "README":
            ws["A1"] = "Supplementary Tables: AlphaGenome T2D liver variant-to-function study"
            ws["A1"].font = Font(name="Arial", size=12, bold=True)
            ws["A2"] = ("Scripts and raw data: https://github.com/ManuHCI/T2D-liver-AlphaGenome  "
                        "(scripts 01-15 regenerate every table).")
            ws["A2"].font = Font(name="Arial", size=9, italic=True)
        else:
            ws["A1"] = caps[ws.title]
            ws["A1"].font = Font(name="Arial", size=10, bold=True)
            if ws.title in blocks:
                row = 3
                for title, d in blocks[ws.title]:
                    ws.cell(row=row, column=1).font = Font(name="Arial", size=10, bold=True, color="2A78D6")
                    for c in ws[row + 1]:
                        if c.value is not None:
                            c.font = Font(name="Arial", size=9, bold=True)
                            c.fill = PatternFill("solid", fgColor="EEF2F8")
                    row += len(d) + 4
            else:
                for c in ws[3]:
                    c.font = Font(name="Arial", size=9, bold=True)
                    c.fill = PatternFill("solid", fgColor="EEF2F8")
                ws.freeze_panes = "A4"
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
        ws.column_dimensions["A"].width = 26 if ws.title != "README" else 30
        if ws.title == "README":
            ws.column_dimensions["B"].width = 120
            for r in ws.iter_rows(min_row=5):
                for c in r:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT)
    print("saved", OUT, "| sheets:", [ws.title for ws in wb.worksheets])


if __name__ == "__main__":
    main()
